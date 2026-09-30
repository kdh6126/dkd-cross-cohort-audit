#!/usr/bin/env python
"""Extra deck slides: the multi-modal ladder and the independent artifact replication.

Kept in its own module rather than inline in make_ppt.py because these two slides carry long
multi-line Korean strings, and writing them through a shell heredoc repeatedly corrupted the
`\\n` escapes. Importing make_ppt is safe - its main() is guarded by __main__.
"""
from pptx.dml.color import RGBColor

from make_ppt import (add_slide, header, txt, rule, table, picture,
                      INK, MUTED, ACCENT, BLUE)

GREYTXT = RGBColor(0x5A, 0x5A, 0x5A)


def s_ladder(prs, d):
    s = add_slide(prs)
    header(s, 'FRAMING', '왜 멀티모달 오믹스가 필요한가 — 사다리 3단을 실측',
           '"더 많은 모달리티가 더 많은 정보"는 참이지만 반증 불가능한 주장입니다. '
           '단계별로 나눠서 재봤습니다.')
    rungs = [
        ('1단 · 단일 데이터셋', BLUE, [
            '11개 데이터셋을 각각 독립 실행. 상위 50개 feature의 데이터셋 간 평균 일치도 = 0.049.',
            '작은 데이터셋 3개(GSE1009/111154/20602)는 서로는 0.19~0.22로 맞지만 실제 코호트와는 '
            '0.00~0.02 — 일치도가 생물학의 증거는 아닙니다.']),
        ('2단 · 같은 모달리티, 여러 데이터셋', ACCENT, [
            '좋은 절반: 교차코호트 안정성 선택으로 교차폴드 Jaccard 0.03~0.05 → 0.192(3코호트) / '
            '0.303(4코호트), AUROC는 동등 이상.',
            '나쁜 절반: 단순 반복 집계로 뽑힌 20개 중 특이성 게이트 통과는 2개, 5개는 채취 '
            '아티팩트. ADA는 vs 대조 +1.82인데 vs 다른 CKD +0.19 — 공유 교란이 오히려 강화됩니다.']),
        ('3단 · 두 번째 모달리티 — 공개 데이터로는 불가', GREYTXT, [
            'KPMP 프로테오믹스: 당뇨 환자 0명. 2025년 대사체 2건(ST004483/442): 원시 mzXML만 등록. '
            'ST000691: 전부 Progression:Y로 대비군 없음. 혈액 전사체: 신장 데이터와 일치도 0.00~0.075.',
            '→ 같은 DKD 환자에서 두 번째 모달리티를 case/control 설계로 측정한 공개 자원이 없습니다.']),
    ]
    y = 2.05
    for title, color, body in rungs:
        txt(s, 0.6, y, 12.2, 0.32, title, size=15, bold=True, color=color)
        txt(s, 0.85, y + 0.36, 11.9, 1.05, body, size=11, line_spacing=1.28)
        y += 1.60
    rule(s, 0.6, 6.9, 12.2)
    txt(s, 0.6, 7.0, 12.2, 0.4,
        '한 모달리티의 모든 코호트가 공유하는 교란은, 아무리 안정적인 단일모달 방법으로도 '
        '보이지 않습니다.', size=11.5, bold=True)


def s_replication(prs, d):
    s = add_slide(prs)
    header(s, 'REPLICATION', '채취 아티팩트 독립 복제 — 가장 아팠던 한계 해소',
           'GSE175759: 다른 기관, RNA-seq(마이크로어레이 아님), 세뇨관 조직. '
           '비당뇨 CKD 62건 vs 신절제 대조 22건.')
    picture(s, 'F9_artifact_replication.png', 0.6, 2.15, w=10.4)
    rows = [['코호트', '플랫폼', 'IEG 모듈 g', '나머지 유전자'],
            ['ERCB', 'microarray', '-1.23', '~0'],
            ['GSE175759', 'RNA-seq', '-2.26', '+0.69']]
    table(s, 0.6, 6.0, 6.2, rows, col_w=[1.6, 1.6, 1.4, 1.4], size=11, hl_rows=(2,))
    txt(s, 7.1, 6.05, 5.7, 1.2, [
        '독립 코호트에서 효과가 오히려 더 큽니다',
        '(모듈 vs 전사체 전체, Mann-Whitney p = 5.4e-14).',
        '',
        '반면 eGFR 연관 검정은 양성 대조군(FMOD/MMP2/LUM)도 전부 미달 → 이 코호트는 71%가 '
        'IgAN이라 검정 자체가 유효하지 않습니다.'],
        size=10, color=MUTED, line_spacing=1.25)


EXTRA_SLIDES = [s_ladder, s_replication]
