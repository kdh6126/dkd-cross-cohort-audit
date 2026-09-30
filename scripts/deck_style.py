# -*- coding: utf-8 -*-
"""DKD발표.pptx 템플릿의 서식을 코드로 옮긴 것.

템플릿에는 placeholder 가 없고 슬라이드마다 도형을 직접 놓습니다. 그래서 좌표와 글꼴을
여기에 모아 두고, 장표 만드는 쪽은 의미만 다루게 합니다. 값은 템플릿 슬라이드에서 그대로
읽은 것입니다 — 눈대중으로 맞추지 않았습니다.

한글 글꼴은 font.name 만 정하면 라틴 문자에만 붙습니다. 한글은 <a:ea> 를 따로 넣어야
적용되므로 kfont() 로 둘을 함께 설정합니다. 이것을 빠뜨리면 화면에서는 맞아 보이는데
다른 PC 에서 열면 기본 글꼴로 떨어집니다.
"""
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ---- 템플릿에서 읽은 글꼴
F_KICK = '페이퍼로지 7 Bold'      # 눈썹줄
F_TITLE = '페이퍼로지 5 Medium'   # 제목
F_SUB = '페이퍼로지 7 Bold'       # 부제
F_BODY = 'KoPubDotum Medium'      # 본문 · 표 · 패널

# ---- 템플릿에서 읽은 색
BLUE = RGBColor(0x00, 0x66, 0xA5)      # 눈썹줄
INK = RGBColor(0x14, 0x14, 0x14)       # 제목 · 본문
MUTED = RGBColor(0x60, 0x60, 0x60)     # 부제
PANEL_H = RGBColor(0x1E, 0x4B, 0x8C)   # 패널 머리
PANEL_B = RGBColor(0x1F, 0x3A, 0x68)   # 패널 본문
ACCENT = RGBColor(0xB2, 0x5A, 0x0B)    # 강조 (주황)
COVER = RGBColor(0x00, 0x5A, 0x92)
LINE = RGBColor(0xD5, 0xDC, 0xE3)
FILL = RGBColor(0xF4, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x3C, 0x3C, 0x3C)

# ---- 템플릿에서 읽은 좌표 (인치)
KICK_XY = (1.63, 0.34, 6.00, 0.22)
TITLE_XY = (1.63, 0.58, 11.20, 0.50)
SUB_XY = (1.63, 1.10, 11.20, 0.40)
BODY_X, BODY_Y, BODY_W = 0.77, 1.42, 11.00
COL_L, COL_R, COL_W = 0.35, 6.85, 6.20

LAY_COVER, LAY_DIVIDER, LAY_KEY, LAY_PLAIN = 0, 3, 4, 5

# 본문 확대율. 템플릿 판형에 맞춰 장표를 만들어 보니 글자가 화면 위쪽 60%만 채우고
# 아래가 비었습니다. 머리글(제목·눈썹줄)은 템플릿 값이라 그대로 두고 본문만 키웁니다.
BODY_SCALE = 1.09
ROW_SCALE = 1.10


def kfont(run, name, size, bold=False, color=None):
    """라틴과 한글에 같은 글꼴을 건다. <a:ea> 를 빠뜨리면 한글이 기본 글꼴로 떨어진다."""
    run.font.name = name
    run.font.size = Pt(size * (1.0 if name in (F_TITLE, F_KICK) else BODY_SCALE))
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set('typeface', name)
    return run


def rich(par, text, name, size, bold=False, color=None):
    """**별표 두 개** 로 감싼 부분을 실제 굵은 글씨로 만든다.

    처음에는 문자열을 그대로 넣었더니 슬라이드에 별표가 그대로 찍혔습니다. 발표 자료에서
    마크다운 기호가 보이면 그 자체로 흠입니다.
    """
    for i, part in enumerate(str(text).split('**')):
        if part == '':
            continue
        kfont(par.add_run(), name, size, bold or (i % 2 == 1), color).text = part
    return par


def blank(prs, layout=LAY_PLAIN):
    return prs.slides.add_slide(prs.slide_layouts[layout])


def clear(prs):
    """템플릿에 들어 있는 예시 슬라이드를 전부 지운다. 레이아웃과 마스터는 남는다."""
    xml = prs.slides._sldIdLst
    for sid in list(xml):
        prs.part.drop_rel(sid.rId)
        xml.remove(sid)


def tb(slide, x, y, w, h, text, font=F_BODY, size=14, bold=False, color=INK,
       align=PP_ALIGN.LEFT, spacing=1.0, anchor=MSO_ANCHOR.TOP, wrap=True):
    """글상자 하나. text 는 문자열이거나 (레벨, 문자열) 목록."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    items = text if isinstance(text, (list, tuple)) else [text]
    for i, it in enumerate(items):
        lvl, s = it if isinstance(it, tuple) else (0, it)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.level = lvl
        p.line_spacing = spacing
        if i:
            p.space_before = Pt(3 if lvl else 8)
        rich(p, s, font, size, bold, color)
    return box


def head(slide, kicker, title, sub=None):
    """모든 내지의 머리. 템플릿 좌표 그대로."""
    tb(slide, *KICK_XY, kicker, font=F_KICK, size=10, bold=True, color=BLUE)
    tb(slide, *TITLE_XY, title, font=F_TITLE, size=24, bold=True, color=INK)
    if sub:
        tb(slide, *SUB_XY, sub, font=F_SUB, size=12.5, color=MUTED)


def bullets(slide, x, y, w, h, items, size=14):
    """레벨 0 은 굵게, 레벨 1 은 한 단 들여쓰고 작게. 템플릿의 위계를 따른다."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        lvl, s = it if isinstance(it, tuple) else (0, it)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = lvl
        p.line_spacing = 1.18
        if i:
            p.space_before = Pt(11 if lvl == 0 else 3)
        if lvl == 0:
            rich(p, s, F_BODY, size, True, INK)
        else:
            rich(p, '·  ' + s, F_BODY, size - 1.5, False, GREY)
    return box


def panel(slide, x, y, w, title, rows, accent=False, size=11, head_size=12, row_h=0.28):
    """둥근 사각형 머리 + 줄 목록. 템플릿의 데이터 패널과 같은 형태."""
    row_h *= ROW_SCALE
    col = ACCENT if accent else PANEL_H
    hd = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(0.39))
    hd.fill.solid(); hd.fill.fore_color.rgb = FILL
    hd.line.color.rgb = col; hd.line.width = Pt(0.75)
    hd.text_frame.word_wrap = True
    p = hd.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    rich(p, title, F_BODY, head_size, False, col)
    yy = y + 0.47
    for r in rows:
        bold, s = r if isinstance(r, tuple) else (False, r)
        if not str(s).strip():          # 빈 줄은 간격만, 불릿을 찍지 않는다
            yy += row_h * 0.5
            continue
        bx = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                    Inches(x), Inches(yy), Inches(w), Inches(row_h))
        bx.fill.background(); bx.line.fill.background()
        tf = bx.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(0.08); tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        pp = tf.paragraphs[0]; pp.alignment = PP_ALIGN.LEFT
        # 앞에 공백을 둔 줄은 하위 항목이므로 불릿을 찍지 않는다
        mark = '' if str(s).startswith(' ') else '•  '
        rich(pp, mark + str(s), F_BODY, size, bold, ACCENT if bold else PANEL_B)
        yy += row_h
    return yy


def grid(slide, x, y, w, rows, widths=None, size=11, hl=(), row_h=0.30, head_h=0.34):
    """표. 템플릿에 표 도형은 없어서, 같은 글꼴과 색으로 선만 그어 맞춘다."""
    n = len(rows[0])
    row_h *= ROW_SCALE
    head_h *= ROW_SCALE
    widths = widths or [1] * n
    tot = float(sum(widths))
    xs, cx = [], x
    for cw in widths:
        xs.append(cx); cx += w * cw / tot
    yy = y
    for ri, row in enumerate(rows):
        h = head_h if ri == 0 else row_h
        if ri == 0 or ri - 1 in hl:
            bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(yy),
                                        Inches(w), Inches(h))
            bg.fill.solid()
            bg.fill.fore_color.rgb = FILL if ri == 0 else RGBColor(0xFD, 0xF3, 0xE7)
            bg.line.fill.background()
        for ci, cell in enumerate(row):
            cw = w * widths[ci] / tot
            box = slide.shapes.add_textbox(Inches(xs[ci]), Inches(yy), Inches(cw), Inches(h))
            tf = box.text_frame; tf.word_wrap = True
            tf.margin_left = Inches(0.07); tf.margin_right = Inches(0.05)
            tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            bold = ri == 0 or (ri - 1) in hl
            col = PANEL_H if ri == 0 else (ACCENT if (ri - 1) in hl else PANEL_B)
            # 셀 안의 줄바꿈은 그대로 살린다. 긴 설명을 넣을 때 끊는 자리를 손으로
            # 정하는 편이, 자동 줄바꿈이 조사 한 글자를 다음 줄로 넘기는 것보다 낫다.
            for li, seg in enumerate(str(cell).split(chr(10))):
                p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
                p.line_spacing = 1.08
                rich(p, seg, F_BODY, size, bold, col)
        ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(yy + h),
                                    Inches(w), Inches(0.008))
        ln.fill.solid(); ln.fill.fore_color.rgb = LINE if ri else PANEL_H
        ln.line.fill.background()
        yy += h
    return yy


MONO = 'Consolas'


def code(slide, x, y, w, h, lines, size=9.5, title=None):
    """실제 스크립트를 장표에 그대로 싣는다.

    발표에서 "어떻게 구현했나" 를 물으면 말로 설명하는 것보다 코드를 보여 주는 편이 빠름.
    저장소의 파일에서 잘라 온 것이므로 손으로 다시 쓰지 않음.
    """
    bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor(0xF7, 0xF9, 0xFB)
    bg.line.color.rgb = LINE; bg.line.width = Pt(0.75)
    bg.text_frame.word_wrap = True
    tf = bg.text_frame
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    first = True
    if title:
        p = tf.paragraphs[0]; first = False
        p.alignment = PP_ALIGN.LEFT
        kfont(p.add_run(), MONO, size + 0.5, True, PANEL_H).text = title
    for ln in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        # 둥근 사각형의 기본 정렬이 가운데라 첫 줄만 가운데로 붙는다. 코드는 전부 왼쪽.
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.06
        txt = str(ln)
        col = MUTED if txt.lstrip().startswith('#') else INK
        kfont(p.add_run(), MONO, size, False, col).text = txt
    return bg


def note(slide, x, y, w, text, size=10.5, color=MUTED):
    return tb(slide, x, y, w, 0.34, text, size=size, color=color, spacing=1.15)


def picture(slide, path, x, y, w=None, h=None):
    import os
    if not os.path.exists(path):
        return None
    kw = {}
    if w: kw['width'] = Inches(w)
    if h: kw['height'] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)
