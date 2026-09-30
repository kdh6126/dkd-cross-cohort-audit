#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Build a Korean project report for the DKD multi-omics evidence framework.

The report deliberately distinguishes patient-matched multi-omics from the
cross-modal corroboration available in the public resources used here.
"""
from __future__ import annotations

import os
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from matplotlib.patches import FancyBboxPatch


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "reports" / "DKD_멀티오믹스_바이오마커_보고서"
OUT_DOCX = OUT_DIR / "DKD_멀티오믹스_바이오마커_발굴_보고서.docx"
ASSET_DIR = OUT_DIR / "assets"
FONT = "Malgun Gothic"
INK = "17324D"
BLUE = "1677A8"
TEAL = "148A83"
ORANGE = "D45500"
GREY = "667085"


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    props.append(node)


def set_cell_text(cell, value, bold=False, color=None, size=8.7):
    cell.text = str(value)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            set_font(run, size=size, bold=bold, color=color)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def make_framework_figure(path: Path):
    plt.rcParams["font.family"] = FONT
    fig, ax = plt.subplots(figsize=(11.5, 5.7), dpi=180)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 6)
    ax.axis("off")
    ax.text(0.2, 5.65, "DKD 멀티오믹스 증거 프레임워크", fontsize=18,
            fontweight="bold", color="#" + INK)
    ax.text(0.2, 5.25, "환자 매칭 통합분석이 아니라, 각 분자층의 역할과 한계를 구분한 증거 연결", fontsize=10.5,
            color="#" + GREY)

    boxes = [
        (0.3, 3.0, 2.45, 1.4, "전사체\n후보 발굴", "7개 코호트 · 258명\n9,900 유전자", BLUE),
        (3.25, 3.0, 2.45, 1.4, "교란 통제\n후보 선별", "조달 민감도 · 질환 고정 대조\n우선 후보: MOXD1", TEAL),
        (6.2, 3.0, 2.45, 1.4, "단백체\n직교 보강", "환자 비매칭 2개 자료\n12개 후보 측정", ORANGE),
        (9.15, 3.0, 2.45, 1.4, "대사체·유전체\n경계 설정", "경로 수준 탐색\n유전적 과장 방지", "6A5A8A"),
    ]
    for x, y, w, h, title, desc, color in boxes:
        rect = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.04,rounding_size=0.12",
                              linewidth=1.8, edgecolor="#" + color, facecolor="#" + color + "18")
        ax.add_patch(rect)
        ax.text(x + w / 2, y + .93, title, ha="center", va="center", fontsize=12,
                fontweight="bold", color="#" + INK)
        ax.text(x + w / 2, y + .36, desc, ha="center", va="center", fontsize=8.8, color="#" + GREY)
    for x in [2.78, 5.73, 8.68]:
        ax.annotate("", xy=(x + .38, 3.7), xytext=(x, 3.7),
                    arrowprops=dict(arrowstyle="->", lw=1.7, color="#" + GREY))

    notes = [
        (0.5, 1.55, "1. Integrative multi-omics", "동일 참여자 또는 짝지은 시료의 두 층을 직접 연결\n현재 공개 DKD 자원에서는 구성하지 못함", "B54708"),
        (4.15, 1.55, "2. Cross-modal corroboration", "환자 비매칭·유사한 질환 대비의 다른 층에서\n전사체 후보를 독립적으로 점검", TEAL),
        (7.8, 1.55, "3. Within-layer discovery", "충분한 코호트·표본·탐색공간이 있을 때만 발굴\n전사체 가능, 단백체는 현재 보강 역할", BLUE),
    ]
    for x, y, title, desc, color in notes:
        ax.text(x, y + .75, title, fontsize=10.2, fontweight="bold", color="#" + color)
        ax.text(x, y, desc, fontsize=8.7, color="#" + GREY, linespacing=1.55)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_proteome_figure(path: Path, meta: pd.Series):
    plt.rcParams["font.family"] = FONT
    fig, axes = plt.subplots(1, 2, figsize=(11.5, 4.7), dpi=180, gridspec_kw={"width_ratios": [1.25, 1]})
    assays = ["SOMAscan", "LC-MS/MS"]
    candidate = [4 / 7 * 100, 2 / 7 * 100]
    background = [28.27, 10.93]
    x = np.arange(2)
    w = .32
    axes[0].bar(x - w / 2, background, w, label="측정된 비후보 단백질", color="#C9C3BB")
    axes[0].bar(x + w / 2, candidate, w, label="전사체 후보", color="#" + ORANGE)
    for i, val in enumerate(background):
        axes[0].text(i - w / 2, val + 2, f"{val:.0f}%", ha="center", fontsize=10, color="#" + GREY)
    for i, val in enumerate(candidate):
        axes[0].text(i + w / 2, val + 2, f"{val:.0f}%", ha="center", fontsize=10, color="#" + ORANGE)
    axes[0].set_xticks(x, assays)
    axes[0].set_ylim(0, 72)
    axes[0].set_ylabel("q<0.05 도달 비율 (%)")
    axes[0].set_title("단백체 층의 후보 보강", fontweight="bold")
    axes[0].legend(frameon=False, fontsize=8)
    axes[0].spines[["top", "right"]].set_visible(False)

    text = (
        "환자 비매칭 단백체 2개 자료\n\n"
        "측정 후보: 12개\n"
        "한 자료 이상 q<0.05: 5개\n"
        "두 기술에서 모두 유의: MMP7 1개\n\n"
        f"MH 공통 승산비 {meta['cmh_or']:.1f}\n"
        f"95% CI {meta['cmh_or_lo']:.1f}–{meta['cmh_or_hi']:.1f}\n"
        f"CMH p={meta['cmh_p']:.3f}; 순열 p={meta['perm_p_sig']:.3f}\n\n"
        "해석: 후보군의 단백질 수준 보강\n"
        "(단백체 단독 발굴 또는 환자매칭 통합분석 아님)"
    )
    axes[1].axis("off")
    axes[1].text(.04, .94, text, va="top", fontsize=11, linespacing=1.65, color="#" + INK,
                 bbox=dict(boxstyle="round,pad=.8", fc="#F5F8FA", ec="#" + BLUE, lw=1.2))
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def build_docx():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    candidate = pd.read_csv(ROOT / "results/candidates_v2/master_candidate_table.tsv", sep="\t")
    literature = pd.read_csv(ROOT / "results/candidates_v2/literature_review.tsv", sep="\t")
    meta = pd.read_csv(ROOT / "results/proteome_meta/summary.tsv", sep="\t").iloc[0]
    coverage = pd.read_csv(ROOT / "results/proteome_meta/candidate_protein_coverage.tsv", sep="\t")
    audit = pd.read_csv(ROOT / "results/proteome_meta/cohort_audit.tsv", sep="\t")
    feasibility = pd.read_csv(ROOT / "results/multiomics_feasibility/summary.tsv", sep="\t").iloc[0]
    moxd1 = candidate[candidate.gene == "MOXD1"].iloc[0]
    mox_lit = literature[literature.gene == "MOXD1"].iloc[0]

    framework = ASSET_DIR / "multiomics_framework.png"
    proteome = ASSET_DIR / "proteome_corroboration.png"
    make_framework_figure(framework)
    make_proteome_figure(proteome, meta)

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(10.3)
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("DKD 멀티오믹스 바이오마커 발굴 보고서 | ")
        set_font(run, size=8.5, color=GREY)
        add_page_number(footer)

    def heading(text, level=1):
        h = doc.add_heading(text, level=level)
        for run in h.runs:
            set_font(run, size={1: 15, 2: 12, 3: 10.8}.get(level, 10.3), bold=True, color=INK)
        return h

    def paragraph(text="", bold_lead=None, note=False):
        p = doc.add_paragraph()
        if bold_lead:
            r = p.add_run(bold_lead)
            set_font(r, bold=True, color=BLUE)
        r = p.add_run(text)
        set_font(r, color=GREY if note else None)
        p.paragraph_format.space_after = Pt(6)
        return p

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        set_font(p.add_run(text))
        return p

    def table(headers, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for i, head in enumerate(headers):
            set_cell_text(t.rows[0].cells[i], head, bold=True, color="FFFFFF", size=8.6)
            shade(t.rows[0].cells[i], INK)
        for row in rows:
            cells = t.add_row().cells
            for i, value in enumerate(row):
                set_cell_text(cells[i], value, size=8.4)
                if len(t.rows) % 2 == 1:
                    shade(cells[i], "F4F7F9")
        if widths:
            for row in t.rows:
                for i, width in enumerate(widths):
                    row.cells[i].width = Cm(width)
        doc.add_paragraph()
        return t

    def picture(path, width=16.3, caption=None):
        doc.add_picture(str(path), width=Cm(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(caption), size=8.5, color=GREY)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DKD 멀티오믹스 증거기반\n바이오마커 발굴 보고서")
    set_font(r, size=25, bold=True, color=INK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("전사체 기반 후보 발굴과 교차모달 단백체 보강"), size=14, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    set_font(p.add_run("과제 보고서용 | 2026년 9월"), size=10.5, color=GREY)
    doc.add_page_break()

    heading("요약", 1)
    paragraph("본 과제는 공개 인간 신장 데이터를 활용하여 DKD 바이오마커 후보를 발굴하고, 후보가 기술·분자층·대조군 구성의 변화에도 유지되는지를 점검하는 멀티오믹스 증거 프레임워크를 구축하였다.")
    table(["구분", "주요 성과"], [
        ["바이오마커 후보 1종", "MOXD1: 7개 전사체 코호트 기반의 강한 DKD 효과와 다른 신장질환 대비 효과를 보인 우선 검증 후보"],
        ["핵심 방법", "교차코호트 feature selection, 코호트 구성 민감도 분석, 조달 경로 교란 통제, 질환 고정 대조군 비교"],
        ["멀티오믹스 구성", "전사체 후보 발굴 + 환자 비매칭 단백체 직교 보강 + 대사체 경로 탐색 + 유전체 경계 설정"],
        ["단백체 보강", f"12개 후보 측정, 5개가 한 자료 이상에서 q<0.05, MH OR {meta['cmh_or']:.1f} (95% CI {meta['cmh_or_lo']:.1f}–{meta['cmh_or_hi']:.1f})"],
        ["한계", "동일 환자에서 두 분자층을 연결하는 integrative multi-omics는 공개 DKD 자원에서 구성하지 못함"],
    ], widths=[3.2, 13.0])
    paragraph("본 보고서의 멀티오믹스는 환자매칭 통합모형을 뜻하지 않는다. 서로 다른 분자층에서 얻은 증거의 역할과 한계를 명시적으로 구분하여 DKD 후보의 신뢰도를 단계적으로 높이는 evidence framework를 뜻한다.", note=True)

    heading("1. 과제 목표와 산출물", 1)
    paragraph("목표는 공개 DKD 자원에서 재현 가능한 바이오마커 후보를 찾고, 단일 코호트·단일 기술에서만 강해 보이는 신호를 걸러내는 것이다.")
    table(["세부 목표", "적용 방법", "산출물"], [
        ["후보 발굴", "7개 코호트, 258명, 9,900개 유전자에서 교차코호트 일관성 평가", "DKD 후보군 30개"],
        ["교란 통제", "조달 경로 점수, 질환 고정 대조군, 음성 대조", "조달 관련 신호와 질환 관련 후보 분리"],
        ["우선 후보 선정", "효과크기, 다른 CKD 대비, 단일세포 위치, 문헌 근거, 외부 검증을 종합", "우선 검증 후보 MOXD1 1종"],
        ["교차모달 보강", "SOMAscan 및 LC-MS/MS 단백체 자료의 층화 비교", "후보군 단백질 수준 보강 근거"],
    ], widths=[3.0, 8.0, 5.2])

    heading("2. 멀티오믹스 증거 프레임워크", 1)
    picture(framework, caption="그림 1. 본 과제의 멀티오믹스 증거 연결 구조. 환자매칭 통합분석과 환자 비매칭 교차모달 보강을 구분한다.")
    table(["층위", "본 과제에서의 역할", "현재 가능한 주장"], [
        ["Integrative multi-omics", "동일 참여자 또는 짝지은 시료에서 분자층 직접 연결", "공개 DKD 자원에서는 구성하지 못함"],
        ["Cross-modal corroboration", "환자 비매칭 단백체 자료에서 전사체 후보 점검", "후보군의 단백질 수준 보강"],
        ["Within-layer discovery", "각 분자층 내부에서 독립 후보 발굴", "전사체 가능; 단백체는 표본·공통 공간 제약으로 불가"],
    ], widths=[4.2, 7.2, 4.8])
    paragraph("단백체 두 자료는 대조군이 각각 건강인과 비당뇨 병리 사례로 유사하지만 동일하지 않다. 따라서 단백체 결과는 DKD 후보의 직교 보강으로만 해석하며, 동일 환자에서의 분자층 통합 또는 단백체 단독 바이오마커 발굴로 주장하지 않는다.", note=True)

    heading("3. 바이오마커 후보군 발굴 방법", 1)
    table(["단계", "사용 방법", "검증 목적"], [
        ["데이터 조화", "프로브-유전자 매핑, 9,900개 고정 유전자 공간, 코호트 내 표준화", "플랫폼 차이를 줄이고 같은 특성 공간 유지"],
        ["교차코호트 선택", "Leave-one-dataset-out 기반 안정성 평가", "특정 코호트에만 맞는 후보 배제"],
        ["코호트 민감도", "발굴 코호트 조합을 바꾸어 순위 변동 측정", "방법 효과와 코호트 구성 효과 분리"],
        ["교란 통제", "조달 처리 점수 잔차화 및 surrogate-variable 조정", "채취·처리 관련 신호의 후보 진입 억제"],
        ["질환 특이성", "DKD 대 다른 신장질환 대비 및 외부 스트레스 테스트", "단순 신장 손상 신호와 DKD 관련 신호 구분"],
    ], widths=[3.0, 7.5, 5.7])

    heading("4. 바이오마커 후보 1종: MOXD1", 1)
    paragraph("MOXD1은 본 과제의 우선 검증 후보로 선정하였다. 이 후보는 신규성만으로 선택된 것이 아니라, 다수 코호트의 전사체 효과와 질환 고정 대조군에서의 유지 여부를 우선으로 평가한 결과이다.")
    table(["근거 축", "MOXD1 결과", "해석"], [
        ["전사체 DKD 효과", f"Hedges g = {moxd1.g_DKD_vs_control:.2f}", "DKD 대 대조군에서 큰 효과크기"],
        ["다른 CKD 대비", f"Hedges g = {moxd1.g_DKD_vs_otherCKD:.2f}", "다른 신장질환 대비에서도 방향과 효과 유지"],
        ["교란 여부", "donor-driven = False", "공여자/조달 경로가 주된 설명이라는 증거 없음"],
        ["세포 수준 위치", str(moxd1.sn_top_celltype).replace("(<i>", "(").replace("</i>", ""), "신장 조직 내 세포 유형 맥락 제시"],
        ["문헌 상태", f"DKD 관련 문헌 {int(mox_lit.n_dkd)}편, 바이오마커 문헌 0편", "완전 신규는 아니나 DKD 바이오마커 근거는 희소"],
        ["단백체 상태", "DKD 대조 단백체 자료에서 미측정", "단백체 검증이 필요한 후속 후보"],
    ], widths=[3.5, 5.2, 7.5])
    paragraph("결론: MOXD1은 ‘확정 바이오마커’가 아니라 후속 단백질 정량, 독립 임상 코호트, 예후 및 조직특이성 검증으로 발전시켜야 할 우선 후보이다.", bold_lead="해석: ")

    heading("5. 후보군 근거와 등급", 1)
    tier_rows = []
    for tier, count in candidate.groupby("tier").size().items():
        label = {"2 - strong, already reported": "강한 데이터 근거·기보고",
                 "3 - novel, but weak in data": "문헌 희소·데이터 근거 제한",
                 "4 - reported and weak": "기보고·데이터 근거 제한"}.get(tier, tier)
        tier_rows.append([label, int(count), "후속 검증 우선순위는 데이터 근거와 독립 검증 가능성으로 결정"])
    table(["후보군 등급", "개수", "과제 내 해석"], tier_rows, widths=[5.2, 2.0, 9.0])
    picture(ROOT / "results/figures/P7_candidates.png", width=16.0,
            caption="그림 2. 30개 전사체 후보의 증거 행렬. 채워진 칸은 해당 근거가 확인된 경우를 뜻한다.")

    heading("6. 단백체 교차모달 보강", 1)
    table(["자료", "기술", "대상", "후보 측정", "해석상 주의"], [
        ["Mendeley 83k89shdx5", "SOMAscan aptamer", "DKD 23 / 건강인 10", "7개", "MMP7은 원 출처 논문의 기보고 결과"],
        ["PRIDE PXD041884", "Label-free LC-MS/MS", "DKD 5 / 비당뇨 7", "7개", "양 군 모두 FFPE explant이나 전분석 시간은 미보고"],
    ], widths=[3.4, 3.4, 3.4, 2.2, 5.8])
    picture(proteome, caption="그림 3. 전사체 후보군의 환자 비매칭 단백체 교차모달 보강. 층화 결합 결과는 단백체 단독 발굴이 아니라 후보군 보강 근거이다.")
    table(["단백체 결과", "값"], [
        ["두 자료에서 측정된 서로 다른 후보", f"{int(meta.n_candidates_measured)}개"],
        ["한 자료 이상에서 q<0.05", f"{int(meta.n_sig_union)}개"],
        ["두 기술 모두에서 유의·방향 일치", "MMP7 1개"],
        ["층화 결합 효과", f"MH 공통 승산비 {meta.cmh_or:.1f} (95% CI {meta.cmh_or_lo:.1f}–{meta.cmh_or_hi:.1f})"],
        ["민감도 분석", f"검출성 매칭 순열 p={meta.perm_p_matched:.3f}"],
    ], widths=[7.0, 10.0])

    heading("7. 멀티오믹스 자원 현황과 활용 범위", 1)
    table(["분자층", "현재 자원", "과제 내 활용", "제약"], [
        ["전사체", "7개 DKD 코호트, 258명", "후보 발굴 및 교란 진단", "공개 코호트의 조달 비대칭"],
        ["단백체", "환자 비매칭 2개 자료", "후보군 직교 보강", "대조군 정의가 다르고 탐색 공간이 좁음"],
        ["대사체", "ST003255 질환 고정 대조", "경로 수준 탐색 및 후속 확장", "현재 단일 코호트라 독립 발굴 주장 보류"],
        ["유전체", "한국 DKD GWAS", "유전적 과장 방지·경계 설정", "후보군의 강한 enrichment 없음"],
        ["공간 오믹스", "KPMP 공간 대사체/지질체", "향후 통합 가능성 확인", "공개 API의 정량 행렬 접근 제약"],
    ], widths=[2.4, 4.2, 5.4, 5.0])
    paragraph(f"KPMP에서 전사체-단백체는 환자 {int(feasibility.tx_prot_case)}명이 겹치지만 대조군이 0명이며, 전사체-대사체는 환자 {int(feasibility.tx_metab_case)}명과 대조군 {int(feasibility.tx_metab_control)}명이 겹친다. 후자는 공간 영상 데이터 형식 때문에 참여자별 정량 행렬로 바로 결합할 수 없다.", note=True)

    heading("8. 성과 및 후속 계획", 1)
    table(["구분", "성과", "후속 계획"], [
        ["바이오마커 후보", "MOXD1 우선 검증 후보 1종 선정", "독립 임상 코호트에서 단백질 정량 및 예후 연관성 검증"],
        ["방법론", "교란·코호트 구성 민감도를 포함한 후보 발굴 파이프라인", "다른 질환과 추가 코호트로 일반화 평가"],
        ["멀티오믹스", "분자층별 증거 연결 및 한계 관리 체계 구축", "동일 환자 기반 RNA-단백질 또는 RNA-대사체 자료 확보 시 통합모형 적용"],
        ["데이터 자원", "공개 데이터 기반 재현 가능한 분석·검증 체계", "저장소/DOI 공개 후 산출물 재생성 지원"],
    ], widths=[3.2, 7.0, 6.8])
    paragraph("본 과제의 차별점은 여러 오믹스 자료를 단순 병렬 제시하는 데 있지 않다. 각 분자층이 할 수 있는 주장과 할 수 없는 주장을 분리하고, 전사체 후보 발굴–단백체 보강–대사체 경로 탐색–유전체 경계 설정을 하나의 검증 흐름으로 연결한 데 있다.", bold_lead="과제 관점의 핵심 기여: ")

    heading("참고 데이터 및 문헌", 1)
    for item in [
        "공개 전사체 코호트: NCBI GEO 및 KPMP Atlas.",
        "단백체 보강: Hirohama et al., JASN 2023 (Mendeley Data 83k89shdx5); Schwab et al., Proteomics Clinical Applications 2024 (PRIDE PXD041884).",
        "대사체 탐색: Metabolomics Workbench ST003255.",
        "유전체 경계 설정: Korean DKD GWAS summary statistics.",
        "MOXD1 문헌 근거: PMID 34908186.",
    ]:
        bullet(item)
    paragraph("수치 출처: results/candidates_v2/, results/proteome_meta/, results/multiomics_feasibility/의 재생성 가능 산출물. 보고서 생성일 기준 분석 결과를 반영함.", note=True)

    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build_docx()
